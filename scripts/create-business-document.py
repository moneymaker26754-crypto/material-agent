"""Generate the business-focused brief from its Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root / 'docs' / 'business-project-brief.md'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(19), Mm(21)
section.left_margin = section.right_margin = Mm(20)
section.footer_distance = Mm(9)
for name, size in [('Normal', 11), ('Title', 22), ('Heading 1', 15), ('Heading 2', 12), ('List Bullet', 11)]:
    style = document.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.paragraph_format.line_spacing = 1.14
    style.paragraph_format.space_after = Pt(5.5)
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
document.styles['Heading 1'].paragraph_format.space_before = Pt(12)
document.styles['Heading 2'].paragraph_format.space_before = Pt(9)
document.styles['List Bullet'].paragraph_format.left_indent = Mm(5)
document.styles['List Bullet'].paragraph_format.first_line_indent = Mm(-3)

def add_inline(paragraph, text):
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not token:
            continue
        run = paragraph.add_run(token[2:-2] if token.startswith('**') else token[1:-1] if token.startswith('`') else token)
        if token.startswith('**'):
            run.bold = True
        if token.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(10)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('Material Agent  ')
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
document.core_properties.title = 'Material Agent 业务问题与技术方案'
document.core_properties.subject = '物料治理与采购辅助的业务需求和技术实现'
document.core_properties.author = 'Material Agent'

lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [cell.strip() for cell in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', cell) for cell in cells):
                rows.append(cells)
            i += 1
        i -= 1
        table = document.add_table(rows=0, cols=3)
        table.style = 'Table Grid'
        table.autofit = False
        widths = [Mm(53), Mm(45), Mm(72)]
        for column, width in zip(table.columns, widths):
            column.width = width
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement('w:' + edge)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                border.set(qn('w:' + key), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
        for row_index, values in enumerate(rows):
            row = table.add_row()
            cant_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(cant_split)
            if row_index == 0:
                repeat = OxmlElement('w:tblHeader')
                row._tr.get_or_add_trPr().append(repeat)
            for column_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for side in ['top', 'left', 'bottom', 'right']:
                    margin = OxmlElement('w:' + side)
                    margin.set(qn('w:w'), '95')
                    margin.set(qn('w:type'), 'dxa')
                    margins.append(margin)
                props.append(margins)
                if row_index == 0:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'E9ECEF')
                    props.append(shading)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                if column_index == 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, value)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = row_index == 0
        document.add_paragraph().paragraph_format.space_after = Pt(1)
    else:
        style = 'Title' if line.startswith('# ') else 'Heading 1' if line.startswith('## ') else 'Heading 2' if line.startswith('### ') else 'List Bullet' if line.startswith('- ') else 'Normal'
        content = re.sub(r'^(#{1,3} |[-] )', '', line)
        paragraph = document.add_paragraph(style=style)
        add_inline(paragraph, content)
        if line.startswith('## 二 '):
            paragraph.paragraph_format.page_break_before = True
        if line.startswith('整体链路为：'):
            paragraph.paragraph_format.keep_with_next = True
    i += 1

output = source.with_suffix('.docx')
document.save(output)
print(output)
