"""Build the Word project overview from its version-controlled Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root / 'docs' / 'project-overview.md'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
    style = document.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.size = Pt(11 if name in ['Normal', 'List Bullet'] else 24 if name == 'Title' else 16 if name == 'Heading 1' else 12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
document.core_properties.title = 'Material Agent 项目梳理'
document.core_properties.subject = '物料治理、采购辅助、Agent 工程架构与验证'
document.core_properties.author = 'Material Agent'
lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line.startswith('```'):
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            p = document.add_paragraph(lines[i]); p.paragraph_format.space_after = Pt(2)
            for run in p.runs: run.font.name = 'Consolas'; run.font.size = Pt(9)
            i += 1
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            row = [cell.strip() for cell in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', cell) for cell in row): rows.append(row)
            i += 1
        i -= 1
        table = document.add_table(rows=1, cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        for row_index, row in enumerate(rows):
            cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = text
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.space_before = Pt(4)
                    for run in paragraph.runs: run.font.size = Pt(10); run.font.color.rgb = RGBColor(0, 0, 0)
            if row_index == 0:
                repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for cell in cells:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E9ECEF'); cell._tc.get_or_add_tcPr().append(shade)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs: run.bold = True
        document.add_paragraph().paragraph_format.space_after = Pt(1)
    elif line.startswith('# '): document.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '): document.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '): document.add_paragraph(line[4:], 'Heading 2')
    elif line.startswith('- '): document.add_paragraph(line[2:], 'List Bullet')
    else: document.add_paragraph(line)
    i += 1
output = root / 'docs' / 'project-overview.docx'
document.save(output)
print(output)
